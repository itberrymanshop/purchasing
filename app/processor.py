from pathlib import Path
from datetime import datetime
import calendar
import json
import math
import re
import openpyxl

MONTHS = {"jan": 1, "januari": 1, "feb": 2, "februari": 2, "mar": 3, "maret": 3, "apr": 4, "april": 4, "mei": 5, "may": 5, "jun": 6, "juni": 6, "jul": 7, "juli": 7, "agu": 8, "agt": 8, "agustus": 8, "sep": 9, "september": 9, "okt": 10, "oktober": 10, "nov": 11, "november": 11, "des": 12, "desember": 12}
SHEET_NAMES = {
    "stok": ("stok", "ketersediaan stok penjualan"), "periode": ("penjualan_periode", "penjualan barang per periode"), "harga": ("master_harga", "master harga jual"),
    "total": ("penjualan_total", "penjualan per barang"), "impor": ("impor", "barang masuk impor"), "ppn": ("ppn", "produk pm np"), "discontinue": ("discontinue", "produk discontinue"), "pareto": ("pareto", "produk pareto"),
}

def clean(value): return re.sub(r"\s+", " ", str(value or "").strip()).lower()
def number(value):
    if value in (None, ""): return 0.0
    if isinstance(value, (int, float)): return float(value)
    value = str(value).replace("Rp", "").replace(".", "").replace(",", ".").strip()
    try: return float(value)
    except ValueError: return 0.0

def header_index(ws, row=1):
    return {clean(cell.value): index for index, cell in enumerate(ws[row]) if clean(cell.value)}
def cell_by_alias(values, headers, *aliases):
    for alias in aliases:
        for header, index in headers.items():
            if header == alias or alias in header:
                return values[index] if index < len(values) else None
    return None

def sheet(wb, key):
    aliases = SHEET_NAMES[key]
    return next((candidate for candidate in wb.worksheets if clean(candidate.title) in aliases), None)

def rows(ws, header_row=1, data_row=None):
    if not ws: return
    headers = header_index(ws, header_row)
    for values in ws.iter_rows(min_row=data_row or header_row + 1, values_only=True):
        if any(value not in (None, "") for value in values): yield values, headers

def code_of(values, headers):
    value = cell_by_alias(values, headers, "kode barang", "kode #", "sku", "marketplace - sku", "grosir - sku", "mp & gr - sku")
    return str(value or "").strip()

def month_from_header(text):
    normalized = clean(text)
    return next((number for name, number in MONTHS.items() if re.search(rf"\b{name}\b", normalized)), None)

def merged_value(ws, row, column):
    cell = ws.cell(row, column)
    if cell.value not in (None, ""):
        return cell.value
    for merged in ws.merged_cells.ranges:
        if merged.min_row <= row <= merged.max_row and merged.min_col <= column <= merged.max_col:
            return ws.cell(merged.min_row, merged.min_col).value
    return cell.value

def normalized_row(ws, row_number):
    return tuple(merged_value(ws, row_number, column) for column in range(1, ws.max_column + 1))

def period_sales(ws):
    if not ws: return {}, []
    current_month = datetime.now().month
    fallback_months = [((current_month - 4 + offset) % 12) + 1 for offset in range(4)]
    branches = ("grosir", "klevo", "marketplace", "subtotal")
    month_order = []
    for group in range(4):
        header = merged_value(ws, 1, 3 + group * 4)
        month_order.append(month_from_header(header) or fallback_months[group])
    output = {}
    first_data_row = None
    for row_number in range(2, ws.max_row + 1):
        value = merged_value(ws, row_number, 1)
        if value not in (None, ""):
            first_data_row = row_number
            break
    if first_data_row is None: return {}, month_order
    for row_number in range(first_data_row, ws.max_row + 1):
        values = normalized_row(ws, row_number)
        code = str(values[0] or "").strip()
        if not code: continue
        item = output.setdefault(code, {})
        for group_index, month in enumerate(month_order):
            start = 2 + group_index * 4
            data = {branch: number(values[start + branch_index]) if start + branch_index < len(values) else 0 for branch_index, branch in enumerate(branches[:3])}
            data["subtotal"] = sum(data.values())
            item[month] = data
    return output, month_order

def keyed_values(ws, value_aliases):
    result = {}
    for values, headers in rows(ws):
        code = code_of(values, headers)
        if code: result[code] = cell_by_alias(values, headers, *value_aliases)
    return result

def process_workbook(path, rules=None):
    rules = rules or {"lead_time_import": 80, "lead_time_lokal": 15, "umur_berryman_import": 90, "umur_klevo_import": 90, "umur_berryman_lokal": 30}
    wb = openpyxl.load_workbook(path, data_only=True)
    stock_ws = sheet(wb, "stok")
    if not stock_ws: raise ValueError("Sheet 'ketersediaan stok penjualan' tidak ditemukan.")
    sales, month_order = period_sales(sheet(wb, "periode"))
    if not month_order: raise ValueError("Header bulan pada sheet penjualan barang per periode tidak ditemukan.")
    current = max(month_order)
    previous, two_previous = current - 1, current - 2
    year = datetime.now().year
    days_in_current = calendar.monthrange(year, current)[1]
    days_elapsed = min(datetime.now().day, days_in_current) if current == datetime.now().month else days_in_current
    prices = {code: number(value) for code, value in keyed_values(sheet(wb, "harga"), ("grosir min", "harga grosir")).items()}
    totals = {code: number(value) for code, value in keyed_values(sheet(wb, "total"), ("penjualan", "total penjualan", "omzet")).items()}
    import_details = {code: str(value or "").strip() for code, value in keyed_values(sheet(wb, "impor"), ("keterangan",)).items()}
    import_codes = set(import_details)
    tax = {}
    for code, value in keyed_values(sheet(wb, "ppn"), ("pm / np", "pm np")).items():
        category = str(value or "").upper().strip()
        category = re.sub(r"\s+", " ", category)
        tax[code] = "PM" if category == "PM" else "NP" if category == "NP" else category
    discontinued = {code: str(value or "AKTIF").upper().strip() or "AKTIF" for code, value in keyed_values(sheet(wb, "discontinue"), ("keterangan barang", "keterangan")).items()}
    pareto_types = {}
    pareto_sheet = sheet(wb, "pareto")
    if pareto_sheet:
        for values, headers in rows(pareto_sheet):
            marketplace_sku = str(cell_by_alias(values, headers, "marketplace sku") or "").strip()
            grosir_sku = str(cell_by_alias(values, headers, "grosir sku") or "").strip()
            mp_gr_sku = str(cell_by_alias(values, headers, "mp gr sku") or "").strip()
            if marketplace_sku and marketplace_sku not in pareto_types:
                pareto_types[marketplace_sku] = "MARKETPLACE"
            if grosir_sku and grosir_sku not in pareto_types:
                pareto_types[grosir_sku] = "GROSIR"
            if mp_gr_sku:
                pareto_types[mp_gr_sku] = "MP_GR"
    result = []
    for position, (values, headers) in enumerate(rows(stock_ws), 1):
        code = code_of(values, headers)
        if not code: continue
        monthly = sales.get(code, {})
        prior_sales = monthly.get(previous, {}).get("subtotal", 0)
        current_sales = monthly.get(current, {}).get("subtotal", 0)
        historical_sales = [monthly.get(month, {}).get("subtotal", 0) for month in month_order[:3]]
        projection = current_sales * days_in_current / max(days_elapsed, 1)
        average_sales = (sum(historical_sales) + projection) / 4
        ads = average_sales / days_in_current if days_in_current else 0
        stock_gudang = number(cell_by_alias(values, headers, "stok gudang"))
        ordered = number(cell_by_alias(values, headers, "dipesan"))
        sold = number(cell_by_alias(values, headers, "dijual"))
        available_raw = cell_by_alias(values, headers, "stok dapat dijual")
        available = number(available_raw) if available_raw not in (None, "") else max(0, stock_gudang - sold)
        previous_daily_sales = prior_sales / days_in_current if days_in_current else 0
        on_hand = available / previous_daily_sales if previous_daily_sales else None
        on_way = (available + ordered) / previous_daily_sales if previous_daily_sales else None
        change = prior_sales - monthly.get(two_previous, {}).get("subtotal", 0)
        kind = "IMPOR" if code in import_codes else "LOKAL"
        condition = discontinued.get(code, "AKTIF")
        lead_time = float(rules["lead_time_import"] if kind == "IMPOR" else rules["lead_time_lokal"])
        is_klevo = str(cell_by_alias(values, headers, "nama barang") or "").strip().upper().startswith("KLEVO")
        brand = "KLEVO" if is_klevo else "BERRYMAN"
        inventory_age = float(rules["umur_klevo_import"] if is_klevo else rules["umur_berryman_import"] if kind == "IMPOR" else rules["umur_berryman_lokal"])
        target_stock = ads * inventory_age if ads else 0
        recommendation = target_stock
        reason = "Target stok menjadi rekomendasi stok."
        if not ads: reason = "Tidak ada penjualan historis; evaluasi manual."
        elif on_way is not None: reason = f"CSOH + OTW {on_way:.1f} hari; umur persediaan {inventory_age:.0f} hari; rekomendasi stok {target_stock:,.0f} unit."
        status = "DISCONTINUE" if condition == "DISCONTINUE" else ("N/A" if on_way is None else "CRITICAL" if on_way < lead_time else "NEED ORDER" if on_way < inventory_age else "SUFFICIENT")
        result.append({
            "no": position, "sku": code, "supplier": str(cell_by_alias(values, headers, "nama pemasok", "pemasok utama") or ""),
            "name": str(cell_by_alias(values, headers, "nama barang") or ""), "cbm": number(cell_by_alias(values, headers, "cbm")),
            "condition": condition, "kind": kind, "tax": tax.get(code, ""), "import_detail": import_details.get(code, ""), "pareto": code in pareto_types, "pareto_type": pareto_types.get(code, ""),
            "stock_gudang": stock_gudang, "ordered": ordered, "sold": sold, "available": available,
            "months": monthly, "month_order": month_order, "current_month": current, "previous_sales": prior_sales,
            "projection": projection, "change": change,
            "trend": "NAIK" if change > 0 else "TURUN" if change < 0 else "STABIL", "price": prices.get(code, 0),
            "total_sales": totals.get(code, 0), "on_hand": on_hand, "on_way": on_way, "status": status,
            "lead_time": lead_time, "brand": brand, "inventory_age": inventory_age, "ads": ads, "target_stock": target_stock, "restock_qty": recommendation, "restock_reason": reason,
            "calculation": {"historical_sales": historical_sales, "average_sales": average_sales, "projection": projection, "previous_sales": prior_sales, "current_sales": current_sales, "sales_days": days_in_current, "sales_days_elapsed": days_elapsed, "previous_daily_sales": previous_daily_sales, "available": available, "ordered": ordered, "inventory_age": inventory_age},
        })
    return result

def save_report(upload_dir, filename, records):
    active, history = upload_dir / "active_report.json", upload_dir / "audit_reports.json"
    old, audit = (json.loads(active.read_text()) if active.exists() else None), (json.loads(history.read_text()) if history.exists() else [])
    if old: audit.insert(0, old)
    report = {"filename": filename, "processed_at": datetime.now().isoformat(timespec="seconds"), "records": records}
    active.write_text(json.dumps(report, ensure_ascii=False)); history.write_text(json.dumps(audit[:30], ensure_ascii=False)); return report

def load_active(upload_dir):
    path = upload_dir / "active_report.json"; return json.loads(path.read_text()) if path.exists() else None
def load_audit(upload_dir):
    path = upload_dir / "audit_reports.json"; return json.loads(path.read_text()) if path.exists() else []
